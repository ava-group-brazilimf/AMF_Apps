"""
Generate AVA Fabric Agents Inventory Excel file.

Scans the repository filesystem directly (agents/, templates/, workflows/,
scripts, schemas, configs, checklists, references) instead of relying on a
hand-maintained list — so agents that exist on disk but are not (yet) wired
into a module.yaml manifest still show up in the inventory.

Phases follow src/modules/ava-fabric-agents/master-orchestrator/agents/master-orchestrator.md
(the authoritative pipeline definition), not the older per-module comments in
the root module.yaml, which are stale (e.g. still list tech-stack as "F3").
"""
import re
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ─────────────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent
MODULES_ROOT = REPO_ROOT / "src" / "modules" / "ava-fabric-agents"
SHARED_ROOT = REPO_ROOT / "src" / "shared"

OUTPUT_PATH = SCRIPT_DIR / "AVA-FABRIC-INVENTORY.xlsx"

# ── Phases (source of truth: master-orchestrator.md § Agent Team) ────────────
PHASE_MAP = {
    "master-orchestrator": "F0 — Master Orchestrator",
    "asis-diagnostic": "F1 — AS-IS Diagnostic",
    "tobe-architecture": "F2 — TO-BE Architecture",
    "prototype": "F3 — Prototype",
    "tech-stack": "F4 — Stack / Codegen",
    "qa-agents": "F5 — QA",
    "devops-agents": "F6 — DevOps",
    "deliverables": "F7 — Deliverables",
    "summary": "Cross-Phase — Summary",
    "shared": "Shared",
}
PHASE_ORDER = [
    "F0 — Master Orchestrator",
    "F1 — AS-IS Diagnostic",
    "F2 — TO-BE Architecture",
    "F3 — Prototype",
    "F4 — Stack / Codegen",
    "F5 — QA",
    "F6 — DevOps",
    "F7 — Deliverables",
    "Cross-Phase — Summary",
    "Shared",
    "Root",
]

EXCLUDE_NAME_PARTS = ("__pycache__",)
EXCLUDE_SUFFIXES = (".backup", ".pyc")
EXCLUDE_NAME_SUBSTR = (".pre-patch",)


def module_of(path: Path) -> str:
    """First path component under MODULES_ROOT, or 'shared'/'root' otherwise."""
    try:
        rel = path.relative_to(MODULES_ROOT)
        return rel.parts[0]
    except ValueError:
        pass
    try:
        path.relative_to(SHARED_ROOT)
        return "shared"
    except ValueError:
        return "root"


def phase_of(path: Path) -> str:
    mod = module_of(path)
    if mod == "root":
        return "Root"
    return PHASE_MAP.get(mod, "Shared")


def is_excluded(path: Path) -> bool:
    if any(part in EXCLUDE_NAME_PARTS for part in path.parts):
        return True
    if path.suffix in EXCLUDE_SUFFIXES:
        return True
    if any(sub in path.name for sub in EXCLUDE_NAME_SUBSTR):
        return True
    return False


def read_frontmatter(path: Path):
    """Return (frontmatter_dict, body) for a markdown file with YAML frontmatter."""
    try:
        text = path.read_text(encoding="utf-8")
    except (UnicodeDecodeError, OSError):
        return {}, ""
    if not text.startswith("---"):
        return {}, text
    parts = text.split("---", 2)
    if len(parts) < 3:
        return {}, text
    try:
        fm = yaml.safe_load(parts[1]) or {}
        if not isinstance(fm, dict):
            fm = {}
    except yaml.YAMLError:
        fm = {}
        m = re.search(r"^name:\s*(.+)$", parts[1], re.MULTILINE)
        if m:
            fm["name"] = m.group(1).strip().strip('"')
    return fm, parts[2]


def clean_description(desc, fallback=""):
    if not desc:
        return fallback
    desc = str(desc)
    for marker in ("Ativa com:", "Ativa quando", "Routing key:", "Tracked in:", "Invocado ", "v1.", "v2."):
        idx = desc.find(marker)
        if idx > 20:
            desc = desc[:idx]
    lines = [l.strip(" -") for l in desc.splitlines() if l.strip()]
    text = " ".join(lines).strip()
    if not text:
        return fallback
    if len(text) > 220:
        text = text[:217].rsplit(" ", 1)[0] + "..."
    return text


def is_stub(fm: dict) -> bool:
    version = str(fm.get("version", ""))
    return "stub" in version.lower()


def is_deprecated(fm: dict) -> bool:
    if fm.get("deprecated") is True:
        return True
    return "deprecated" in str(fm.get("version", "")).lower()


def build_subagent_registry():
    """Parse module.yaml nested `skills:` lists (the authoritative structural
    signal for a sub-agent — e.g. ava-asis-db-analyzer's 4 DB skills, or
    ava-asis-security-orchestrator's 7 security skills) into {abs_path: parent_id}."""
    registry = {}
    for module_yaml in MODULES_ROOT.glob("*/module.yaml"):
        try:
            data = yaml.safe_load(module_yaml.read_text(encoding="utf-8")) or {}
        except yaml.YAMLError:
            continue
        module_dir = module_yaml.parent
        for agent in (data.get("agents") or []):
            parent_id = agent.get("id", "")
            for sub in (agent.get("skills") or []):
                sub_file = sub.get("file")
                if sub_file:
                    registry[(module_dir / sub_file).resolve()] = parent_id
    return registry


SUBAGENT_REGISTRY = build_subagent_registry()

# Self-declared-in-prose fallbacks, for sub-agents not (yet) wired into a module.yaml
# nested `skills:` block (e.g. ava-stack-build-fixer, invoked only by build-validator).
INVOKED_EXCLUSIVELY_PATTERN = re.compile(r"invocad[oa]\s+exclusivamente\s+pel[oa]s?\s+([\w-]+)", re.IGNORECASE)
SELF_DECLARE_PATTERN = re.compile(r"sub-?(?:agent[e]?|skill)\s+d[oe]\s+([\w-]+)", re.IGNORECASE)


def detect_subagent(path: Path, name: str, raw_desc: str):
    """Return (is_subagent, parent_id_or_empty)."""
    parent = SUBAGENT_REGISTRY.get(path.resolve())
    if parent:
        return True, parent
    m = INVOKED_EXCLUSIVELY_PATTERN.search(raw_desc)
    if m:
        return True, m.group(1)
    # Generic self-declaration ("Sub-agent do X") — skip agents that are themselves
    # orchestrators (they legitimately describe managing N sub-agents, e.g.
    # "Security orchestrator — 7 sub-agents...", which is not a self-declaration).
    if "orchestrator" not in name.lower() and SELF_DECLARE_PATTERN.search(raw_desc):
        m = SELF_DECLARE_PATTERN.search(raw_desc)
        candidate = m.group(1)
        return True, (candidate if len(candidate) > 2 else "")
    return False, ""


def title_from_heading(path: Path, default: str) -> str:
    try:
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("# "):
                return line.strip().lstrip("#").strip()
    except (UnicodeDecodeError, OSError):
        pass
    return default


def name_from_filename(path: Path) -> str:
    return path.stem.replace("_", " ").replace("-", " ").title()


# ── Scanners ──────────────────────────────────────────────────────────────────

def scan_agents():
    rows = []
    for path in sorted(MODULES_ROOT.glob("**/agents/**/*.md")):
        if is_excluded(path):
            continue
        fm, _ = read_frontmatter(path)
        if is_deprecated(fm):
            continue
        name = fm.get("name") or path.stem
        raw_desc = str(fm.get("description") or "")
        desc = clean_description(raw_desc, fallback=f"Agent spec: {path.stem}")
        status = "Planning" if is_stub(fm) else "Done"
        is_sub, parent = detect_subagent(path, name, raw_desc)
        if is_sub:
            asset_type = "Sub-Agent"
            if parent and parent not in desc:
                desc = f"{desc} (sub-agent of {parent})"
        else:
            asset_type = "Agent"
        rows.append((status, name, desc, asset_type, phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_templates():
    rows = []
    globs = list(MODULES_ROOT.glob("**/templates/**/*")) + list(SHARED_ROOT.glob("templates/**/*"))
    for path in sorted(set(globs)):
        if not path.is_file() or is_excluded(path):
            continue
        if path.suffix in (".js",):
            continue
        name = path.name
        desc = f"Template — {path.relative_to(REPO_ROOT).as_posix()}"
        rows.append(("Done", name, desc, "Template", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_scripts():
    rows = []
    globs = list(MODULES_ROOT.glob("**/*.py")) + list(SHARED_ROOT.glob("**/*.py"))
    for path in sorted(set(globs)):
        if is_excluded(path) or path.name == "__init__.py":
            continue
        desc = f"Python script — {path.relative_to(REPO_ROOT).as_posix()}"
        status = "Done"
        if path.name.startswith("test_"):
            desc = f"Test script — {path.relative_to(REPO_ROOT).as_posix()}"
        rows.append((status, path.name, desc, "Script", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_workflows():
    rows = []
    for path in sorted(MODULES_ROOT.glob("**/workflows/**/workflow.md")):
        if is_excluded(path):
            continue
        fm, _ = read_frontmatter(path)
        name = fm.get("name") or path.parent.name
        desc = clean_description(fm.get("description"), fallback=f"Workflow: {name}")
        rows.append(("Done", name, desc, "Workflow", phase_of(path), str(path.relative_to(REPO_ROOT))))

    for path in sorted(MODULES_ROOT.glob("**/workflows/**/step-*.md")):
        if is_excluded(path):
            continue
        workflow_name = path.parent.parent.name if path.parent.name == "steps" else path.parent.name
        title = title_from_heading(path, default=path.stem)
        name = f"{path.stem} ({workflow_name})"
        desc = f"{title} — step of workflow '{workflow_name}'"
        rows.append(("Done", name, desc, "Workflow Step", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_checklists():
    rows = []
    globs = list(MODULES_ROOT.glob("**/checklists/*.md")) + list(SHARED_ROOT.glob("checklists/*.md"))
    for path in sorted(set(globs)):
        if is_excluded(path):
            continue
        title = title_from_heading(path, default=name_from_filename(path))
        rows.append(("Done", path.stem, f"Checklist — {title}", "Checklist", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_schemas():
    rows = []
    for path in sorted(SHARED_ROOT.glob("schemas/*.json")):
        if is_excluded(path):
            continue
        rows.append(("Done", path.name, f"JSON Schema — {path.stem.replace('-', ' ').replace('.', ' ')}", "Schema", "Shared", str(path.relative_to(REPO_ROOT))))
    return rows


def scan_configs():
    rows = []
    config_paths = [REPO_ROOT / "module.yaml"] + list(MODULES_ROOT.glob("*/module.yaml"))
    for path in sorted(set(config_paths)):
        if not path.exists() or is_excluded(path):
            continue
        label = "module.yaml (root)" if path.parent == REPO_ROOT else f"module.yaml ({path.parent.name})"
        rows.append(("Done", label, f"Module config — {path.relative_to(REPO_ROOT).as_posix()}", "Config", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def scan_references():
    rows = []
    globs = (
        list(SHARED_ROOT.glob("data/**/*.yaml"))
        + list(SHARED_ROOT.glob("data/**/*.md"))
        + list(SHARED_ROOT.glob("data/**/*.version"))
        + list(MODULES_ROOT.glob("*/shared/*.md"))
        + list(MODULES_ROOT.glob("shared/*.md"))
        + list(MODULES_ROOT.glob("*/data/*.yaml"))
    )
    for path in sorted(set(globs)):
        if not path.is_file() or is_excluded(path):
            continue
        rows.append(("Done", path.name, f"Reference/shared doc — {path.relative_to(REPO_ROOT).as_posix()}", "Reference", phase_of(path), str(path.relative_to(REPO_ROOT))))
    return rows


def build_inventory():
    rows = []
    rows += scan_agents()
    rows += scan_templates()
    rows += scan_scripts()
    rows += scan_workflows()
    rows += scan_checklists()
    rows += scan_schemas()
    rows += scan_configs()
    rows += scan_references()
    # Dedup on (name, type, path) just in case a glob matched twice
    seen = set()
    deduped = []
    for row in rows:
        key = (row[1], row[3], row[5])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(row)
    # Sort by Phase order, then Type, then Name
    def sort_key(row):
        status, name, desc, atype, phase, relpath = row
        phase_idx = PHASE_ORDER.index(phase) if phase in PHASE_ORDER else len(PHASE_ORDER)
        return (phase_idx, atype, name.lower())
    deduped.sort(key=sort_key)
    return deduped


inventory = build_inventory()

# ── Styles ──────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="FF4B0082", end_color="FF4B0082", fill_type="solid")  # Avanade purple
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Segoe UI", size=10)
BOLD_FONT = Font(name="Segoe UI", bold=True, size=10)
TITLE_FONT = Font(name="Segoe UI", bold=True, size=14, color="4B0082")
SUBTITLE_FONT = Font(name="Segoe UI", bold=True, size=11, color="666666")

STATUS_COLORS = {
    "Done":        PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid"),
    "In Progress": PatternFill(start_color="FFFFFFCC", end_color="FFFFFFCC", fill_type="solid"),
    "New":         PatternFill(start_color="FFBDD7EE", end_color="FFBDD7EE", fill_type="solid"),
    "Planning":    PatternFill(start_color="FFE2EFDA", end_color="FFE2EFDA", fill_type="solid"),
}

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

wb = openpyxl.Workbook()

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 1: Inventory
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Inventory"

ws.merge_cells("A1:G1")
ws["A1"] = "AVA Fabric Agents — Asset Inventory"
ws["A1"].font = TITLE_FONT
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:G2")
ws["A2"] = f"Generated: {date.today().isoformat()} | Total Assets: {len(inventory)} | Source: filesystem scan of src/modules/ava-fabric-agents + src/shared"
ws["A2"].font = SUBTITLE_FONT
ws.row_dimensions[2].height = 20

headers = ["#", "Status", "Asset Name", "Description", "Type", "Phase", "Path"]
col_widths = [5, 12, 40, 70, 14, 24, 60]

for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=4, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[4].height = 22

for i, (status, name, desc, asset_type, phase, relpath) in enumerate(inventory, 1):
    row = i + 4
    values = [i, status, name, desc, asset_type, phase, relpath]
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=(col_idx == 4))
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    status_cell = ws.cell(row=row, column=2)
    if status in STATUS_COLORS:
        status_cell.fill = STATUS_COLORS[status]
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 18

ws.auto_filter.ref = f"A4:G{len(inventory) + 4}"
ws.freeze_panes = "A5"

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 2: Executive Summary
# ═══════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Executive Summary")

ws2.merge_cells("A1:D1")
ws2["A1"] = "AVA Fabric Agents — Executive Summary"
ws2["A1"].font = TITLE_FONT
ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[1].height = 30

ws2.merge_cells("A2:D2")
ws2["A2"] = f"Generated: {date.today().isoformat()}"
ws2["A2"].font = SUBTITLE_FONT

# ── Summary by Type (Type | Count) ──
ws2["A4"] = "Summary by Asset Type"
ws2["A4"].font = Font(name="Segoe UI", bold=True, size=12, color="4B0082")
ws2.merge_cells("A4:D4")

type_headers = ["Type", "Count", "% of Total"]
for col_idx, header in enumerate(type_headers, 1):
    cell = ws2.cell(row=5, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER

type_counts = Counter(item[3] for item in inventory)
type_order = ["Agent", "Sub-Agent", "Workflow", "Workflow Step", "Template", "Script", "Schema", "Config", "Checklist", "Reference"]

total = len(inventory)
row = 6
for t in type_order:
    if t in type_counts:
        count = type_counts[t]
        pct = f"{count / total * 100:.1f}%"
        for col_idx, val in enumerate([t, count, pct], 1):
            cell = ws2.cell(row=row, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        row += 1

for col_idx, val in enumerate(["TOTAL", total, "100%"], 1):
    cell = ws2.cell(row=row, column=col_idx, value=val)
    cell.font = BOLD_FONT
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
    cell.fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")

row += 2

# ── Agents vs Sub-Agents ──
ws2.cell(row=row, column=1, value="Agents vs Sub-Agents").font = Font(name="Segoe UI", bold=True, size=12, color="4B0082")
ws2.merge_cells(f"A{row}:D{row}")
row += 1

for col_idx, header in enumerate(["Category", "Count", "% of Total"], 1):
    cell = ws2.cell(row=row, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
row += 1

top_level_agents = sum(1 for i in inventory if i[3] == "Agent")
sub_agents = sum(1 for i in inventory if i[3] == "Sub-Agent")
agent_family_total = top_level_agents + sub_agents
for label, count in [
    ("Total Agentes (top-level, dispatched by an orchestrator/module.yaml)", top_level_agents),
    ("Total Sub-Agentes (managed internally by another agent)", sub_agents),
    ("TOTAL (Agentes + Sub-Agentes)", agent_family_total),
]:
    is_total_row = label.startswith("TOTAL")
    pct = f"{count / agent_family_total * 100:.1f}%" if agent_family_total else "0.0%"
    for col_idx, val in enumerate([label, count, pct], 1):
        cell = ws2.cell(row=row, column=col_idx, value=val)
        cell.font = BOLD_FONT if is_total_row else DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        if is_total_row:
            cell.fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    row += 1

row += 1

# ── Summary by Phase ──
ws2.cell(row=row, column=1, value="Summary by Phase (per master-orchestrator.md)").font = Font(name="Segoe UI", bold=True, size=12, color="4B0082")
ws2.merge_cells(f"A{row}:E{row}")
row += 1

phase_headers = ["Phase", "Agents", "Sub-Agents", "Other Assets", "Total"]
for col_idx, header in enumerate(phase_headers, 1):
    cell = ws2.cell(row=row, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
row += 1

for phase in PHASE_ORDER:
    agents = sum(1 for item in inventory if item[4] == phase and item[3] == "Agent")
    subagents = sum(1 for item in inventory if item[4] == phase and item[3] == "Sub-Agent")
    others = sum(1 for item in inventory if item[4] == phase and item[3] not in ("Agent", "Sub-Agent"))
    phase_total = agents + subagents + others
    if phase_total > 0:
        for col_idx, val in enumerate([phase, agents, subagents, others, phase_total], 1):
            cell = ws2.cell(row=row, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        row += 1

for col_idx, val in enumerate(["TOTAL",
                                top_level_agents,
                                sub_agents,
                                sum(1 for i in inventory if i[3] not in ("Agent", "Sub-Agent")),
                                total], 1):
    cell = ws2.cell(row=row, column=col_idx, value=val)
    cell.font = BOLD_FONT
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
    cell.fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")

ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 14
ws2.column_dimensions["D"].width = 16
ws2.column_dimensions["E"].width = 14

# ── Summary by Status ──
row += 2
ws2.cell(row=row, column=1, value="Summary by Status").font = Font(name="Segoe UI", bold=True, size=12, color="4B0082")
ws2.merge_cells(f"A{row}:D{row}")
row += 1

for col_idx, header in enumerate(["Status", "Count", "% of Total"], 1):
    cell = ws2.cell(row=row, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
row += 1

status_counts = Counter(item[0] for item in inventory)
for status in ["Done", "In Progress", "New", "Planning"]:
    count = status_counts.get(status, 0)
    if count > 0:
        pct = f"{count / total * 100:.1f}%"
        for col_idx, val in enumerate([status, count, pct], 1):
            cell = ws2.cell(row=row, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
            if col_idx == 1 and status in STATUS_COLORS:
                cell.fill = STATUS_COLORS[status]
        row += 1

# Save
wb.save(OUTPUT_PATH)
print(f"Excel saved to: {OUTPUT_PATH}")
print(f"Total assets: {total}")
print(f"Types: {dict(type_counts)}")
print(f"Phases: {dict(Counter(item[4] for item in inventory))}")
